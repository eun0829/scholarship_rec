# -*- coding: utf-8 -*-
"""[COMMON] XLSX -> 텍스트 (openpyxl). 시트별 상위 행만 사용 (대부분 양식/명단이라 가볍게 처리)."""
from __future__ import annotations

from .results import ExtractError, ExtractResult

MAX_ROWS = 200


def extract(path, settings=None, **_) -> ExtractResult:
    try:
        from openpyxl import load_workbook
        wb = load_workbook(str(path), read_only=True, data_only=True)
    except Exception as e:
        raise ExtractError("XLSX 열기 실패: {}".format(e))
    try:
        return _read(wb)
    finally:
        wb.close()                       # read_only 모드는 파일 핸들을 계속 잡고 있음


def _read(wb) -> ExtractResult:
    parts = []
    for ws in wb.worksheets:
        rows = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= MAX_ROWS:
                rows.append("... (이하 생략)")
                break
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                rows.append(" | ".join(cells))
        if rows:
            parts.append("## 시트: {}\n{}".format(ws.title, "\n".join(rows)))
    text = "\n\n".join(parts).strip()
    if not text:
        raise ExtractError("XLSX에서 추출된 내용 없음")
    return ExtractResult(text=text, method="TEXT")
