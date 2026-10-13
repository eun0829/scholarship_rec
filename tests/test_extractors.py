# -*- coding: utf-8 -*-
import io
import zipfile
import zlib

import pytest

from scholarship.common.extractors import ExtractError, extract
from scholarship.common.extractors import image_, zip_
from scholarship.common.http import sniff_ok


# ---------------------------------------------------------------- HWP (실제 샘플)
def test_hwp_samples_extract_conditions(settings, sample_hwps):
    found = {}
    for f in sample_hwps:
        r = extract(f, "hwp", settings)
        assert r.method == "TEXT" and len(r.text) > 1000
        found[f.name] = r.text
    joined = "\n".join(found.values())
    # 표 안 조건이 살아 있어야 함 (기본 hwp5txt 는 '<표>' 로만 나옴)
    assert "<표>" not in joined
    assert any("4.5 만점 기준 2.8 이상" in t for t in found.values())          # 인천: 표 안 성적 기준
    assert any("[글상자]" in t and "신청 방법&순서" in t for t in found.values())  # 울산: 글상자


def test_hwp_broken_file(settings, tmp_path):
    p = tmp_path / "broken.hwp"
    p.write_bytes(b"not a hwp file")
    with pytest.raises(ExtractError):
        extract(p, "hwp", settings)


# ---------------------------------------------------------------- HWPX (합성 샘플 - 실제 파일로 재검증 필요)
def test_hwpx_synthetic(settings, tmp_path):
    xml = ('<hs:sec xmlns:hs="http://www.hancom.co.kr/hwpml/2011/section" '
           'xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph">'
           '<hp:p><hp:run><hp:t>신청기간: 2026. 9. 1.~9. 30.</hp:t></hp:run></hp:p>'
           '<hp:p><hp:run><hp:tbl><hp:tr>'
           '<hp:tc><hp:subList><hp:p><hp:run><hp:t>구분</hp:t></hp:run></hp:p></hp:subList></hp:tc>'
           '<hp:tc><hp:subList><hp:p><hp:run><hp:t>자격</hp:t></hp:run></hp:p></hp:subList></hp:tc>'
           '</hp:tr></hp:tbl></hp:run></hp:p></hs:sec>')
    p = tmp_path / "a.hwpx"
    with zipfile.ZipFile(p, "w") as z:
        z.writestr("Contents/section0.xml", xml)
    r = extract(p, "hwpx", settings)
    assert "신청기간: 2026. 9. 1.~9. 30." in r.text and "구분 | 자격" in r.text


def test_hwpx_without_sections(settings, tmp_path):
    p = tmp_path / "e.hwpx"
    with zipfile.ZipFile(p, "w") as z:
        z.writestr("mimetype", "x")
    with pytest.raises(ExtractError):
        extract(p, "hwpx", settings)


# ---------------------------------------------------------------- DOCX / XLSX
def make_docx(path):
    from docx import Document
    d = Document()
    d.add_paragraph("2026학년도 장학생 모집 공고")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "자격", "평점 3.0/4.5 이상"
    a = t.cell(1, 0).merge(t.cell(1, 1))
    a.text = "소득 8분위 이내"
    d.save(str(path))


def test_docx_with_table_and_merged_cell(settings, tmp_path):
    p = tmp_path / "a.docx"
    make_docx(p)
    r = extract(p, "docx", settings)
    assert "장학생 모집 공고" in r.text and "자격 | 평점 3.0/4.5 이상" in r.text
    assert r.text.count("소득 8분위 이내") == 1                  # 병합 셀 중복 없음


def test_xlsx(settings, tmp_path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "양식"
    ws.append(["성명", "학번"])
    ws.append([None, None])
    p = tmp_path / "a.xlsx"
    wb.save(str(p))
    assert "성명 | 학번" in extract(p, "xlsx", settings).text


# ---------------------------------------------------------------- PDF
def _pdf(path, text=None, scanned=False, pages=1):
    import pymupdf
    doc = pymupdf.open()
    for i in range(pages):
        page = doc.new_page()
        if text:
            page.insert_text((72, 100), text, fontname="helv", fontsize=12)
        if scanned:
            from PIL import Image
            buf = io.BytesIO()
            Image.new("RGB", (400, 300), (200, 200, 200)).save(buf, "PNG")
            page.insert_image(pymupdf.Rect(50, 50, 450, 350), stream=buf.getvalue())
    doc.save(str(path))


def test_pdf_text_type(settings, tmp_path):
    p = tmp_path / "t.pdf"
    _pdf(p, text="Scholarship notice. Deadline 2026-10-25. GPA 3.0 or higher required for all applicants.")
    r = extract(p, "pdf", settings)
    assert r.method == "TEXT" and "Deadline 2026-10-25" in r.text and r.pages == 1


def test_pdf_scanned_goes_to_vision(settings, tmp_path):
    settings.max_pdf_pages_vision = 2
    p = tmp_path / "s.pdf"
    _pdf(p, scanned=True, pages=4)
    r = extract(p, "pdf", settings)
    assert r.method == "VISION" and len(r.images) == 2 and r.images[0].media_type == "image/png"
    assert any("앞 2쪽" in w for w in r.warnings)


# ---------------------------------------------------------------- 이미지
def _png(path, w, h):
    from PIL import Image
    Image.new("RGB", (w, h), (255, 255, 255)).save(str(path), "PNG")


def test_image_normal_and_tall(settings, tmp_path):
    p = tmp_path / "a.png"
    _png(p, 1300, 1839)                                           # 샘플 포스터 크기 -> 분할하지 않음
    r = extract(p, "png", settings)
    assert r.method == "VISION" and len(r.images) == 1
    p2 = tmp_path / "tall.png"
    _png(p2, 900, 5000)                                           # 매우 긴 이미지 -> 분할
    r2 = extract(p2, "png", settings)
    assert 2 <= len(r2.images) <= image_.MAX_TILES and "분할" in r2.images[0].label


def test_image_corrupt_and_tiny(settings, tmp_path):
    p = tmp_path / "bad.jpg"
    p.write_bytes(b"\xff\xd8garbage")
    with pytest.raises(ExtractError):
        extract(p, "jpg", settings)
    t = tmp_path / "tiny.png"
    _png(t, 5, 5)
    with pytest.raises(ExtractError):
        extract(t, "png", settings)


# ---------------------------------------------------------------- ZIP
def test_zip_members_roles_and_limits(settings, tmp_path):
    d = tmp_path / "x.docx"
    make_docx(d)
    z = tmp_path / "a.zip"
    with zipfile.ZipFile(z, "w") as zf:
        zf.write(d, "공고문.docx")
        zf.writestr("합격자 명단.xlsx", b"PK-not-used")          # 제외 대상 -> 처리하지 않음
        zf.writestr("readme.exe", b"MZ")                          # 미지원 -> 경고
    r = extract(z, "zip", settings)
    assert "[ZIP 내 파일: 공고문.docx" in r.text and "명단" not in r.text
    assert any("제외" in w for w in r.warnings)


def test_zip_too_many_and_nested(settings, tmp_path):
    settings.max_zip_members = 2
    z = tmp_path / "many.zip"
    with zipfile.ZipFile(z, "w") as zf:
        for i in range(3):
            zf.writestr("f{}.txt".format(i), "x")
    with pytest.raises(ExtractError):
        extract(z, "zip", settings)
    with pytest.raises(ExtractError):
        extract(z, "zip", settings, depth=1)                        # 중첩 ZIP 금지


def test_zip_legacy_korean_filename_decoding():
    info = zipfile.ZipInfo("x")
    info.filename = "공고문.hwp".encode("cp949").decode("cp437")      # 옛 압축 프로그램이 만든 이름
    info.flag_bits = 0
    assert zip_._member_name(info) == "공고문.hwp"


def test_unsupported_formats(settings, tmp_path):
    p = tmp_path / "a.doc"
    p.write_bytes(b"x")
    with pytest.raises(ExtractError, match="DOC"):
        extract(p, "doc", settings)


def test_sniff():
    assert sniff_ok("hwp", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1") and not sniff_ok("hwp", b"<html>")
    assert sniff_ok("pdf", b"%PDF-1.7") and not sniff_ok("pdf", b"<!DOCTYPE")
    assert sniff_ok("xyz", b"anything")


# ---------------------------------------------------------------- HWP 변환 실패 진단 / 미리보기 폴백
def test_hwp_flag_reasons_and_stderr_noise_filter():
    from scholarship.common.extractors import hwp
    assert hwp.flag_reasons(0b1) == []                                   # 압축만 -> 정상 문서
    assert hwp.flag_reasons(0b110) == ["암호 설정 문서", "배포용 문서"]
    assert hwp.flag_reasons(0b10000) == ["DRM 보안 문서"]
    err = (b"undefined UnderlineStyle value: 15\ndefined name/values: {...}\n" * 500
           + b"Traceback (most recent call last):\n  File x\nValueError: boom\n")
    tail = hwp.stderr_tail(err)
    assert "ValueError: boom" in tail and "UnderlineStyle" not in tail


def _fail_run(monkeypatch, stderr=b"UnderlineStyle noise\nValueError: boom\n"):
    import subprocess
    from scholarship.common.extractors import hwp
    monkeypatch.setattr(hwp.subprocess, "run", lambda *a, **k: subprocess.CompletedProcess(a, 1, b"", stderr))
    return hwp


def test_hwp_failure_reports_cause_when_no_preview(settings, tmp_path, monkeypatch):
    hwp = _fail_run(monkeypatch)
    p = tmp_path / "x.hwp"
    p.write_bytes(b"not an ole file")
    with pytest.raises(ExtractError) as ei:
        hwp.extract(p, settings)
    msg = str(ei.value)
    assert "HWP 5.0 형식이 아님" in msg and "ValueError: boom" in msg and "UnderlineStyle" not in msg


def test_hwp_failure_falls_back_to_preview_text(settings, sample_hwps, monkeypatch):
    hwp = _fail_run(monkeypatch)
    r = hwp.extract(sample_hwps[0], settings)                           # 실제 HWP: 변환은 막았지만 PrvText 는 읽힘
    assert r.partial is True and r.method == "TEXT" and 50 <= len(r.text) <= 1100
    assert "미리보기 텍스트" in r.warnings[0] and "문서 버전" in r.warnings[0]


def test_hwp_preview_and_header_of_real_samples(sample_hwps):
    from scholarship.common.extractors import hwp
    for f in sample_hwps:
        ver, flags = hwp.read_header(f)
        assert ver.startswith("5.") and hwp.flag_reasons(flags) == []
        assert len(hwp.preview_text(f)) > 100


def test_hwp_environment_check(monkeypatch):
    import builtins
    from scholarship.common.extractors import hwp
    assert hwp.check_environment() == (True, "")
    real_import = builtins.__import__

    def fake_import(name, *a, **k):
        if name == "hwp5.hwp5proc":
            raise ModuleNotFoundError("No module named 'six'", name="six")
        return real_import(name, *a, **k)
    monkeypatch.setattr(builtins, "__import__", fake_import)
    ok, msg = hwp.check_environment()
    assert ok is False and "pip install six" in msg and "'six'" in msg
