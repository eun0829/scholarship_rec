#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""[RECOMMEND] 정확도 자동 점검 (AI 호출 없음, 크레딧 0): 가상 학생 × 저장 공고 × 학교를 판정해 문제 목록과 보고서를 만든다.

  python export_notices.py                              # 먼저 DB 공고를 내려받는다 (한 번)
  python run_accuracy.py                                # 보고서: out_accuracy/accuracy.md, accuracy.xlsx
  python run_accuracy.py --students 400 --seed 7        # 가상 학생 수와 시드
  python run_accuracy.py --school 시흥                  # 학교 하나만 (기본: 5곳 모두)
  python run_accuracy.py --update-baseline              # 지금 결과를 기준선으로 저장 (data/accuracy/baseline.json)
  python run_accuracy.py --strict                       # 오류·불변 규칙 위반이 있으면 종료 코드 1
기준선이 있으면 새로 생긴 문제와 해결된 문제를 같이 보여 준다.
"""
import argparse
import json
import sys
from pathlib import Path

from scholarship.recommend import audit as A
from scholarship.recommend import population as P
from scholarship.recommend.matrix import SCHOOLS

CODE_HELP = {
    "PREFIX_IN_COMMON": "공통 문장에 '유형명: …' 이 남음 (추출 프롬프트·재추출)", "HIGH_SCHOOL_TEXT": "고등학생 전용 문장이 저장됨 (판정에서는 제외됨)",
    "STRUCTURABLE_TEXT": "숫자가 든 조건이 문장으로만 남음 (구조 항목·항목 사전 후보)", "REGION_NO_SIDO": "지역에 시·도가 없음 (재추출)",
    "REGION_NO_AREAS": "지역 요건에 지명이 없음", "GPA_NO_SCALE": "평점 기준에 만점이 없음", "VARIANT_EMPTY": "유형에 구조 조건이 없음 (공고문 확인 필요로만 표시됨)",
    "SPECIAL_OTHER": "필수 신분이 '기타'라 판정 불가 (사전 항목 후보)", "AMOUNT_MISSING": "요약에는 금액이 있는데 amount 가 비어 있음", "ONE_VARIANT": "유형이 1개뿐", "OLD_SCHEMA": "예전 스키마로 추출됨 (v3 재추출 대상)", "NO_DEADLINE": "마감일 없음",
}


def to_markdown(rep, diff=None):
    L = ["# 정확도 점검 보고서", "",
         "- 공고 {}건 × 학생 {}명 × 학교 {} = 판정 {:,}건".format(rep["notices"], rep["students"], rep["school"], rep["pairs"]),
         "- 판정 분포: " + " / ".join("{} {:,}".format(k, v) for k, v in rep["overall"].items()),
         "- 확정률(충족+미달): {:.1%}".format(rep["decided_rate"]),
         "- 오류 {}건 · 불변 규칙 위반 {}건".format(len(rep["crashes"]), len(rep["invariants"])), ""]
    if diff:
        L += ["## 기준선과 비교", "- 확정률: {} → {}".format(*["{:.1%}".format(x) if x is not None else "-" for x in diff["decided_rate"]]),
              "- 불변 규칙 위반: {} → {}".format(*diff["invariants"]),
              "- 새로 생긴 추출 문제: {}건 · 해결된 추출 문제: {}건".format(len(diff["new_lint"]), len(diff["fixed_lint"])),
              "- 새로 생긴 분포 문제: {}건 · 해결된 분포 문제: {}건".format(len(diff["new_flags"]), len(diff["fixed_flags"])), ""]
        for nk, code in diff["new_lint"][:30]:
            L.append("  - 새 문제 {} {}".format(nk, code))
        L.append("")
    L += ["## 1. 조건부로 남는 이유 순위 (고칠 목록)", "",
          "단독 = 그 이유 하나만 풀면 충족이 되는 쌍 수 (고쳤을 때 충족이 늘어나는 정도).", "",
          "| 종류 | 쌍 수 | 단독 | 공고 수 | 예시 | 학생 입력으로 해결 |", "|---|---|---|---|---|---|"]
    struct = [r for r in rep["unknown_ranking"] if not r["kind"].startswith(("other:", "exclusion:"))]
    sent = [r for r in rep["unknown_ranking"] if r["kind"].startswith(("other:", "exclusion:"))]
    L += ["| {kind} | {pairs:,} | {sole:,} | {notices} | {sample} | {student_can_resolve} |".format(**r) for r in struct[:25]]
    L += ["", "### 문장 조건(구조로 옮기지 못한 조건) 순위", "", "| 문장 | 쌍 수 | 단독 | 공고 수 |", "|---|---|---|---|"]
    L += ["| {sample} | {pairs:,} | {sole:,} | {notices} |".format(**r) for r in sorted(sent, key=lambda x: (-x["notices"], -x["sole"], -x["pairs"]))[:30]]
    L += ["", "### 고치면 충족이 가장 많이 늘어나는 순서 (단독 기준)", "", "| 종류 | 단독 | 공고 수 | 예시 |", "|---|---|---|---|"]
    L += ["| {kind} | {sole:,} | {notices} | {sample} |".format(**r) for r in sorted(rep["unknown_ranking"], key=lambda x: -x["sole"])[:15] if r["sole"]]
    L += ["", "## 2. 추출 품질 문제 (공고 수)", "", "| 코드 | 공고 수 | 설명 | 공고 |", "|---|---|---|---|"]
    L += ["| {} | {} | {} | {} |".format(r["code"], r["notices"], CODE_HELP.get(r["code"], ""), ", ".join(r["ids"][:8]) + (" …" if len(r["ids"]) > 8 else "")) for r in rep["lint_summary"]]
    flagged = [p for p in rep["per_notice"] if p["flag"]]
    L += ["", "## 3. 공고 분포 확인 대상 (충족 0명)", "", "| 공고 | 제목 | 충족 | 조건부 | 미달 | 구분 | 가장 흔한 미달 이유 / 충족을 막는 이유 |", "|---|---|---|---|---|---|---|"]
    L += ["| {notice} | {title} | {match} | {cond} | {no} | {flag} | {top_fail} |".format(**p) for p in flagged]
    cv = rep["coverage"]
    L += ["", "## 4. 구조화 커버리지", "", "- 구조로 담긴 조건 {} · 문장으로만 남은 조건 {} → 커버리지 {:.1%}".format(cv["structured"], cv["sentences"], cv["overall"]), "",
          "문장 조건에 자주 나오는 낱말 (사전 항목 후보):", "", "| 낱말 | 공고 수 |", "|---|---|"]
    L += ["| {word} | {notices} |".format(**w) for w in cv["top_words"][:20]]
    L += ["", "커버리지가 낮은 공고:", "", "| 공고 | 제목 | 구조 | 문장 | 커버리지 |", "|---|---|---|---|---|"]
    L += ["| {notice} | {title} | {structured} | {sentences} | {coverage:.0%} |".format(**p) for p in sorted(cv["per_notice"], key=lambda x: x["coverage"])[:15]]
    if rep["invariants"]:
        L += ["", "## 5. 불변 규칙 위반", "", "| 규칙 | 학생 | 공고 | 변화 |", "|---|---|---|---|"]
        L += ["| {rule} | {student} | {notice} | {detail} |".format(**v) for v in rep["invariants"]]
    if rep["crashes"]:
        L += ["", "## 6. 판정 오류", ""] + ["- {student} × {notice}: {error}".format(**c) for c in rep["crashes"][:30]]
    return "\n".join(L) + "\n"


def to_xlsx(rep, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    F, HF = "Malgun Gothic", PatternFill("solid", fgColor="1F3864")

    def sheet(title, headers, rows, widths, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill, c.alignment = Font(name=F, bold=True, color="FFFFFF", size=10), HF, Alignment(horizontal="center", wrap_text=True)
        for r in rows:
            ws.append(list(r))
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font, c.alignment = Font(name=F, size=10), Alignment(vertical="top", wrap_text=True)
        for i, w in enumerate(widths):
            ws.column_dimensions[chr(65 + i)].width = w
        ws.freeze_panes = "A2"
    sheet("요약", ["항목", "값"], [("공고", rep["notices"]), ("학생", rep["students"]), ("학교", rep["school"]), ("판정 쌍", rep["pairs"]),
                              ("확정률", rep["decided_rate"]), ("오류", len(rep["crashes"])), ("불변 규칙 위반", len(rep["invariants"]))] +
          [("판정 " + k, v) for k, v in rep["overall"].items()], [22, 24], True)
    sheet("조건부 이유 순위", ["종류", "쌍 수", "공고 수", "예시", "학생 입력으로 해결"], [(r["kind"], r["pairs"], r["notices"], r["sample"], r["student_can_resolve"]) for r in rep["unknown_ranking"]], [34, 10, 10, 70, 22])
    sheet("추출 품질", ["공고", "코드", "설명", "내용"], [(nk, i["code"], CODE_HELP.get(i["code"], ""), i["detail"]) for nk, items in sorted(rep["lint"].items()) for i in items], [16, 20, 40, 80])
    sheet("공고 분포", ["공고", "제목", "학생×학교", "충족", "조건부", "미달", "구분", "가장 흔한 미달 이유"],
          [(p["notice"], p["title"], p["total"], p["match"], p["cond"], p["no"], p["flag"], p["top_fail"]) for p in rep["per_notice"]], [16, 40, 12, 8, 8, 8, 14, 60])
    sheet("문장 조건 목록", ["공고", "범위", "문장"], [(x["notice"], x["scope"], x["text"]) for x in rep["coverage"]["sentence_list"]], [16, 18, 100])
    sheet("커버리지", ["공고", "제목", "구조 조건", "문장 조건", "커버리지"], [(p["notice"], p["title"], p["structured"], p["sentences"], p["coverage"]) for p in sorted(rep["coverage"]["per_notice"], key=lambda x: x["coverage"])], [16, 44, 10, 10, 10])
    sheet("불변 규칙 위반", ["규칙", "학생", "공고", "변화"], [(v["rule"], v["student"], v["notice"], v["detail"]) for v in rep["invariants"]], [16, 10, 16, 24])
    wb.save(path)


def reextract_commands(lint, codes, chunk=8):
    """추출 문제가 있는 공고를 출처별로 모아 다시 추출하는 명령을 만든다. 외부01 은 run_ext01.py, 드림스폰은 run_ext02.py."""
    ids = {"EXT01": [], "EXT02": []}
    for nk, items in sorted(lint.items()):
        if any(i["code"] in codes for i in items):
            src, _, sid = nk.partition(":")
            if src in ids:
                ids[src].append(sid)
    out = ["# 다시 추출할 공고: 외부01 {}건, 드림스폰 {}건 (AI 호출 {}건). 한 줄씩 실행하세요.".format(len(ids["EXT01"]), len(ids["EXT02"]), len(ids["EXT01"]) + len(ids["EXT02"]))]
    for src, tmpl in (("EXT01", "python run_ext01.py --days 400 --only {} --force"), ("EXT02", "python run_ext02.py --include-closed --only {} --force")):
        for i in range(0, len(ids[src]), chunk):
            out.append(tmpl.format(" ".join(ids[src][i:i + chunk])))
    return "\n".join(out)


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="정확도 자동 점검 (AI 호출 없음)")
    ap.add_argument("--notices", default="data/accuracy/notices", help="export_notices.py 가 만든 폴더")
    ap.add_argument("--students", type=int, default=250, help="조합으로 만드는 가상 학생 수 (대표 학생 14명은 별도로 항상 포함)")
    ap.add_argument("--seed", type=int, default=20261014)
    ap.add_argument("--school", choices=list(SCHOOLS), help="학교 하나만 (기본: 5곳 모두)")
    ap.add_argument("--out", default="out_accuracy")
    ap.add_argument("--baseline", default="data/accuracy/baseline.json")
    ap.add_argument("--update-baseline", action="store_true")
    ap.add_argument("--strict", action="store_true", help="오류·불변 규칙 위반이 있으면 종료 코드 1")
    ap.add_argument("--print-reextract", action="store_true", help="점검은 하지 않고, 추출 문제가 있는 공고를 다시 추출하는 명령(PowerShell)을 출력")
    ap.add_argument("--codes", default="PREFIX_IN_COMMON,OLD_SCHEMA,REGION_NO_SIDO,GPA_NO_SCALE", help="--print-reextract 가 대상으로 삼을 추출 문제 코드 (쉼표)")
    ap.add_argument("--chunk", type=int, default=8, help="명령 한 줄에 넣을 글번호 수")
    a = ap.parse_args(argv)
    notices = A.load_notices(a.notices)
    if not notices:
        print("공고가 없습니다: {} (먼저 python export_notices.py)".format(a.notices), file=sys.stderr)
        return 2
    if a.print_reextract:
        print(reextract_commands(A.lint_all(notices), [c.strip() for c in a.codes.split(",") if c.strip()], a.chunk))
        return 0
    students = P.load_personas() + P.generate(a.students, a.seed)
    rep = A.build_report(notices, students, a.school)
    base_path = Path(a.baseline)
    diff = None
    if base_path.exists():
        diff = A.diff_baseline(rep, json.loads(base_path.read_text(encoding="utf-8")))
    out = Path(a.out)
    out.mkdir(parents=True, exist_ok=True)
    (out / "accuracy.md").write_text(to_markdown(rep, diff), encoding="utf-8")
    (out / "accuracy.json").write_text(json.dumps(rep, ensure_ascii=False, indent=1), encoding="utf-8")
    to_xlsx(rep, str(out / "accuracy.xlsx"))
    if a.update_baseline:
        base_path.parent.mkdir(parents=True, exist_ok=True)
        base_path.write_text(json.dumps(dict({k: rep[k] for k in ("decided_rate", "invariants", "lint", "per_notice")}, coverage=rep["coverage"]["overall"]), ensure_ascii=False, indent=1), encoding="utf-8")
        print("기준선을 저장했습니다: {}".format(base_path))
    print("공고 {} · 학생 {} · 판정 {:,}건 · 확정률 {:.1%} · 오류 {} · 불변 위반 {}".format(
        rep["notices"], rep["students"], rep["pairs"], rep["decided_rate"], len(rep["crashes"]), len(rep["invariants"])))
    print("보고서: {}".format(out / "accuracy.md"))
    return 1 if a.strict and A.severe_count(rep) else 0


if __name__ == "__main__":
    sys.exit(main())
